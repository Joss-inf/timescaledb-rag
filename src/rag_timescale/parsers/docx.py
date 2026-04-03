from __future__ import annotations

from typing import Any

import docx

from rag_timescale.parsers.base import BaseParser, ParsedDocument


class DOCXParser(BaseParser):
    def parse(self, file_path: str | bytes, metadata: dict[str, Any] | None = None) -> ParsedDocument:
        if isinstance(file_path, bytes):
            import io

            doc = docx.Document(io.BytesIO(file_path))
        else:
            doc = docx.Document(file_path)

        text_parts: list[str] = []
        for para in doc.paragraphs:
            if para.style.name.startswith("Heading"):
                level = para.style.name.replace("Heading ", "")
                prefix = "#" * int(level) if level.isdigit() else "#"
                text_parts.append(f"{prefix} {para.text}")
            else:
                text_parts.append(para.text)

        full_text = "\n\n".join(text_parts)

        return ParsedDocument(
            text=full_text,
            metadata=metadata or {},
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def supported_extensions(self) -> set[str]:
        return {".docx"}

    def supported_mime_types(self) -> set[str]:
        return {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
